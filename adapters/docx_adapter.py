"""docx 适配器（设计 §7.1 / 增补设计 §1 振假名）。

- 段落级翻译，保留段落样式；表格单元格逐段；
- w:ruby 结构化解析：基词内联 + {rN} 注音槽；
- 渲染：含注音槽的段落按"纯文本 run + 原生 w:ruby run"序列重建（锚定译词），
  无法锚定时降级为括号注音（已知边界）；无注音段落走原整段套用路径；
- 双语：bi_inter（译段后插入原文段，斜体）、bi_table（生成左右表格文档）。
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from xml.sax.saxutils import escape

from adapters.base import Block, DocumentModel, FormatError, effectively_empty
from adapters.ruby import RUBY_TOKEN_FULL, restore_ruby, split_anchor
from core.segmentation import is_mostly_latin, join_parts, split_long


def _iter_para_chunks(para):
    """按 XML 顺序产出段落文本片段：('text', s) 或 ('ruby', base, rt)。"""
    from docx.oxml.ns import qn
    for r in para._p.findall(qn("w:r")):
        ruby = r.find(qn("w:ruby"))
        if ruby is not None:
            rb = ruby.find(qn("w:rubyBase"))
            rt_el = ruby.find(qn("w:rt"))
            base = "".join(t.text or "" for t in rb.iter(qn("w:t"))) if rb is not None else ""
            rt = "".join(t.text or "" for t in rt_el.iter(qn("w:t"))) if rt_el is not None else ""
            if base or rt:
                yield ("ruby", base, rt)
                continue
        text = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if text:
            yield ("text", text)


def para_flat_text(para) -> tuple[str, list[dict]]:
    """段落 → (含注音槽的平文本, ruby entries)。"""
    parts: list[str] = []
    entries: list[dict] = []
    for chunk in _iter_para_chunks(para):
        if chunk[0] == "text":
            parts.append(chunk[1])
        else:
            _, base, rt = chunk
            token = "{r%d}" % len(entries)
            entries.append({"token": token, "base": base, "rt": rt, "style": "docx"})
            parts.append(base + token)
    return "".join(parts), entries


class DocxAdapter:
    name = "docx"

    def parse(self, path: Path, opts: dict | None = None) -> DocumentModel:
        from docx import Document as open_docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = open_docx(str(path))
        blocks: list[Block] = []
        paras: list = []

        def add_para(para: "Paragraph") -> None:
            try:
                style_name = para.style.name if para.style is not None else ""
            except Exception:
                style_name = ""
            is_heading = style_name.startswith("Heading") or style_name.startswith("标题")
            text, ruby = para_flat_text(para)
            if effectively_empty(text):
                blocks.append(Block(seq=len(blocks), text=text, translatable=False,
                                    meta={"kind": "empty", "el_idx": len(paras)}))
                paras.append(para)
                return
            parts = split_long(text)
            for i, part in enumerate(parts):
                blocks.append(Block(
                    seq=len(blocks), text=part, is_heading=is_heading,
                    meta={"kind": "para", "para": len(paras), "part": i,
                          "ruby": ruby, "el_idx": len(paras)}))
            paras.append(para)

        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                add_para(Paragraph(child, doc))
            elif child.tag == qn("w:tbl"):
                table = Table(child, doc)
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            add_para(para)

        model = DocumentModel(path=path, fmt="docx", blocks=blocks)
        model.skeleton = {"doc": doc, "paras": paras}
        return model

    def _group_final(self, model, first: Block, translations) -> tuple[str, str]:
        group = [b for b in model.blocks
                 if b.meta.get("kind") == "para" and b.meta.get("para") == first.meta["para"]]
        latin = is_mostly_latin(first.text)
        src_final = join_parts([b.text for b in group], latin)
        tgt_final = join_parts([translations.get(b.seq, b.text) for b in group], latin)
        return src_final, tgt_final

    @staticmethod
    def _rpr_xml(para) -> str:
        from docx.oxml.ns import qn
        runs = para.runs
        if runs:
            rpr = runs[0]._r.find(qn("w:rPr"))
            if rpr is not None:
                from lxml import etree
                return etree.tostring(rpr, encoding="unicode")
        return ""

    @staticmethod
    def _plain_run(para, rpr: str, text: str):
        run = para.add_run(text)
        if rpr:
            from docx.oxml import parse_xml
            from docx.oxml.ns import qn
            el = parse_xml(rpr)
            run._r.insert(0, el)
        return run

    @staticmethod
    def _ruby_run(rpr: str, base: str, rt: str):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        return parse_xml(
            f'<w:r {nsdecls("w")}>{rpr}<w:ruby>'
            f'<w:rubyPr><w:rubyAlign w:val="distributeSpace"/></w:rubyPr>'
            f'<w:rt><w:r>{rpr}<w:t xml:space="preserve">{escape(rt)}</w:t></w:r></w:rt>'
            f'<w:rubyBase><w:r>{rpr}<w:t xml:space="preserve">{escape(base)}</w:t></w:r></w:rubyBase>'
            f'</w:ruby></w:r>')

    def _apply_paragraph(self, para, text: str, entries: list[dict],
                         rt_map: dict) -> None:
        """无注音槽 → 整段套用主样式；有 → 纯文本 run + 原生 w:ruby run 序列重建。"""
        runs = para.runs
        rpr = self._rpr_xml(para)
        if not RUBY_TOKEN_FULL.search(text):
            if runs:
                runs[0].text = text
                for r in runs[1:]:
                    r.text = ""
            else:
                para.add_run(text)
            return
        by_token = {e["token"]: e for e in (entries or [])}
        # 清空现有 run 文本（保留首 run 以携带样式基准）
        if runs:
            for r in runs[1:]:
                r.text = ""
            runs[0].text = ""
        pieces = RUBY_TOKEN_FULL.split(text)
        pending: str = ""
        for piece in pieces:
            if RUBY_TOKEN_FULL.fullmatch(piece):
                entry = by_token.get(piece)
                rt = (rt_map or {}).get(piece) or (entry["rt"] if entry else "")
                before, base = split_anchor(pending)
                if entry is None or base is None:
                    pending = pending + f"（{rt}）"
                    continue
                if before:
                    self._plain_run(para, rpr, before)
                para._p.append(self._ruby_run(rpr, base, rt))
                pending = ""
            else:
                pending += piece
        if pending:
            self._plain_run(para, rpr, pending)

    def render(self, out_path: Path, model: DocumentModel,
               translations: dict[int, str], mode: str = "target",
               ruby_maps: dict[int, dict[str, str]] | None = None) -> None:
        from docx import Document as new_docx

        ruby_maps = ruby_maps or {}
        doc = model.skeleton["doc"]
        paras = model.skeleton["paras"]

        if mode == "bi_table":
            out = new_docx()
            table = out.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "原文"
            table.rows[0].cells[1].text = "译文"
            for b in model.blocks:
                if b.meta.get("kind") != "para" or b.meta.get("part", 0) != 0:
                    continue
                src_final, tgt_final = self._group_final(model, b, translations)
                rt_map = {}
                for gb in model.blocks:
                    if gb.meta.get("para") == b.meta["para"] and gb.meta.get("kind") == "para":
                        rt_map.update(ruby_maps.get(gb.seq) or {})
                src_final = restore_ruby(src_final, b.meta.get("ruby"))
                tgt_final = restore_ruby(tgt_final, b.meta.get("ruby"), rt_map,
                                         style_override="paren")
                row = table.add_row()
                row.cells[0].text = src_final
                row.cells[1].text = tgt_final
            out.save(str(out_path))
            return

        if mode not in ("target", "bi_inter"):
            raise FormatError(f"docx 不支持导出模式：{mode}")

        done: set[int] = set()
        for b in model.blocks:
            if b.meta.get("kind") != "para" or b.meta.get("part", 0) != 0:
                continue
            pidx = b.meta["para"]
            if pidx in done:
                continue
            done.add(pidx)
            para = paras[pidx]
            group = [gb for gb in model.blocks
                     if gb.meta.get("kind") == "para" and gb.meta.get("para") == pidx]
            rt_map = {}
            for gb in group:
                rt_map.update(ruby_maps.get(gb.seq) or {})
            src_final, tgt_final = self._group_final(model, b, translations)
            src_copy = deepcopy(para._p) if mode == "bi_inter" else None
            self._apply_paragraph(para, tgt_final, b.meta.get("ruby"), rt_map)
            if src_copy is not None:
                para._p.addnext(src_copy)
                from docx.text.paragraph import Paragraph as P
                src_para = P(src_copy, para._parent)
                src_text = restore_ruby(src_final, b.meta.get("ruby"))
                if src_para.runs:
                    src_para.runs[0].text = src_text
                    src_para.runs[0].italic = True
                    for r in src_para.runs[1:]:
                        r.text = ""

        doc.save(str(out_path))
