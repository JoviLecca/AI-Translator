"""振假名机制的源语言门控（缺陷修复回归）。

**缺陷**：青空文库式 `《》` 是与源语言无关的正则，非日语项目里的书名号被判成
「基词 + 注音」—— `他读了《红楼梦》第三回。` 变成基词 `了` + 注音槽 `{r0}`，
书名号连同书名一起进注音槽：送翻文本错位、书名内容丢失。全角括号、内联
`<ruby>`、`w:ruby` 同理（都不是"源语言无关"的标记）。

**修复**：`adapters.ruby.ruby_enabled(src_lang)` 作为总开关，只有项目源语言是
日语（`ja` 前缀）时才启用振假名识别；六个适配器（txt / md / srt / html / epub /
docx）都从 `opts["src_lang"]` 取值。导入与导出重解析必须传同一个源语言，
否则 seq 无法重放（见 `core/pipeline.py`）。
"""
import zipfile
from pathlib import Path

import pytest

from adapters import get_adapter
from adapters.ruby import ruby_enabled
from core.pipeline import ExportService, ImportService
from core.project import Project

ZH_LINE = "他读了《红楼梦》第三回。"
JA_LINE = "僕は北瀬一廣《かずひろ》だ。"


def _parse(fmt: str, path: Path, **opts):
    """返回 (model, 可译段文本列表)。"""
    model = get_adapter(fmt).parse(path, opts=opts)
    return model, [b.text for b in model.blocks if b.translatable]


# ---------- 总开关本身 ----------

def test_ruby_enabled_only_for_japanese():
    assert ruby_enabled("ja-JP") is True
    assert ruby_enabled("ja") is True
    assert ruby_enabled(" JA-jp ") is True
    assert ruby_enabled("zh-CN") is False
    assert ruby_enabled("en-US") is False
    assert ruby_enabled("") is False
    assert ruby_enabled(None) is False


# ---------- txt ----------

def test_txt_book_title_untouched_for_non_japanese_source(tmp_path):
    p = tmp_path / "zh.txt"
    p.write_text(ZH_LINE + "\n", encoding="utf-8")

    _, zh_texts = _parse("txt", p, src_lang="zh-CN")
    assert zh_texts == [ZH_LINE], zh_texts          # 书名号原样保留
    assert "{r" not in zh_texts[0]

    # 同一个文件在日语源下仍按青空文库式识别（证明差异来自门控而非正则本身）
    _, ja_texts = _parse("txt", p, src_lang="ja-JP")
    assert "{r0}" in ja_texts[0], ja_texts


def test_txt_loose_paren_gate(tmp_path):
    """宽松识别（全角括号）同样受源语言门控。"""
    p = tmp_path / "zh2.txt"
    p.write_text("魔法（まほう）が使える。\n", encoding="utf-8")

    _, zh_texts = _parse("txt", p, ruby_loose=True, src_lang="zh-CN")
    assert zh_texts == ["魔法（まほう）が使える。"], zh_texts

    _, ja_texts = _parse("txt", p, ruby_loose=True, src_lang="ja-JP")
    assert ja_texts == ["魔法{r0}が使える。"], ja_texts


# ---------- md（含内联 <ruby>） ----------

def test_md_book_title_and_inline_ruby_gate(tmp_path):
    p = tmp_path / "zh.md"
    p.write_text(f"# 序章\n\n{ZH_LINE}\n\n他用了<ruby>魔法<rt>まほう</rt></ruby>。\n",
                 encoding="utf-8")

    _, zh_texts = _parse("md", p, src_lang="zh-CN")
    joined = "\n".join(zh_texts)
    assert "{r" not in joined, joined
    assert "# 序章" in joined and ZH_LINE in joined
    assert "<ruby>魔法<rt>まほう</rt></ruby>" in joined, joined   # 标签按普通文本透传

    _, ja_texts = _parse("md", p, src_lang="ja-JP")
    joined_ja = "\n".join(ja_texts)
    assert "{r0}" in joined_ja and "まほう" not in joined_ja, joined_ja


# ---------- srt ----------

def test_srt_cue_with_book_title_gate(tmp_path):
    p = tmp_path / "zh.srt"
    p.write_text("1\n00:00:01,000 --> 00:00:04,000\n" + ZH_LINE + "\n",
                 encoding="utf-8")

    _, zh_texts = _parse("srt", p, src_lang="zh-CN")
    assert zh_texts == [ZH_LINE], zh_texts

    _, ja_texts = _parse("srt", p, src_lang="ja-JP")
    assert "{r0}" in ja_texts[0], ja_texts


# ---------- html / epub（结构化 <ruby>） ----------

HTML_RUBY = "<html><body><p>彼の名は<ruby>一廣<rt>かずひろ</rt></ruby>だ。</p></body></html>"


def test_html_ruby_element_gate(tmp_path):
    p = tmp_path / "r.html"
    p.write_text(HTML_RUBY, encoding="utf-8")

    _, zh_texts = _parse("html", p, src_lang="zh-CN")
    assert "{r" not in zh_texts[0], zh_texts
    assert "一廣" in zh_texts[0]                  # 内容不丢：基词 + 注音按普通文字保留
    assert "かずひろ" in zh_texts[0], zh_texts

    _, ja_texts = _parse("html", p, src_lang="ja-JP")
    assert "一廣{r0}" in ja_texts[0], ja_texts


def _mini_epub(tmp_path: Path, body: str) -> Path:
    """最小可用 EPUB：没有 container.xml/OPF，适配器会退回扫描 .xhtml 扩展名。

    ⚠️ XHTML 必须带编码声明：适配器按 **bytes** 解析内容文档（lxml 不接受带编码
    声明的 str），没有声明时 lxml 会按 windows-1252 猜，UTF-8 正文变乱码。
    """
    p = tmp_path / "m.epub"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("OEBPS/ch1.xhtml",
                   '<?xml version="1.0" encoding="utf-8"?>'
                   '<html xmlns="http://www.w3.org/1999/xhtml"><body><p>'
                   + body + "</p></body></html>")
    return p


def test_epub_ruby_gate(tmp_path):
    ep = _mini_epub(tmp_path, "彼は<ruby>魔法<rt>まほう</rt></ruby>を使った。")

    _, zh_texts = _parse("epub", ep, src_lang="zh-CN")
    assert "{r" not in zh_texts[0], zh_texts

    _, ja_texts = _parse("epub", ep, src_lang="ja-JP")
    assert "魔法{r0}" in ja_texts[0], ja_texts


# ---------- docx（原生 w:ruby） ----------

def _docx_with_ruby(path: Path) -> Path:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("彼の名は")
    para._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:ruby>'
        f'<w:rubyPr><w:rubyAlign w:val="distributeSpace"/></w:rubyPr>'
        f'<w:rt><w:r><w:t>かずひろ</w:t></w:r></w:rt>'
        f'<w:rubyBase><w:r><w:t>一廣</w:t></w:r></w:rubyBase>'
        f'</w:ruby></w:r>'))
    para.add_run("だ。")
    doc.save(str(path))
    return path


def test_docx_wruby_gate(tmp_path):
    src = _docx_with_ruby(tmp_path / "r.docx")

    model, zh_texts = _parse("docx", src, src_lang="zh-CN")
    assert "{r" not in zh_texts[0], zh_texts
    assert zh_texts[0] == "彼の名は一廣（かずひろ）だ。", zh_texts   # 退化为纯文本，不丢内容
    assert not (model.blocks[0].meta.get("ruby") or [])

    _, ja_texts = _parse("docx", src, src_lang="ja-JP")
    assert ja_texts[0] == "彼の名は一廣{r0}だ。", ja_texts


# ---------- 端到端：导入 → 导出重解析 seq 必须一致 ----------

def test_chinese_project_import_export_keeps_book_title(tmp_path):
    """中文项目的书名号必须完整进库、完整出库（导入与导出重解析同门控）。

    修复前：导入时书名号被拆成注音槽（`src_text` 变成「他读了{r0}第三回。」），
    送翻的基词只剩一个 `了`；书名内容在翻译链路上彻底丢失。
    """
    project = Project.create(tmp_path / "p", name="p", src_lang="zh-CN",
                             tgt_lang="en-US", provider_id="mock")
    f = tmp_path / "zh.txt"
    f.write_text(ZH_LINE + "\n\n第二段。\n", encoding="utf-8")
    imported, errs = ImportService(project).import_files([f])
    assert not errs and imported

    src_texts = [s["src_text"] for s in project.db.list_segments(translatable=True)]
    assert ZH_LINE in src_texts, src_texts
    assert all("{r" not in t for t in src_texts), src_texts

    ids = [d["id"] for d in project.db.list_documents()]
    results = ExportService(project).export(ids, mode="target")
    assert all(r["ok"] for r in results), results
    out = Path(results[0]["out"]).read_text(encoding="utf-8")
    assert ZH_LINE in out, out
    assert "{r" not in out, out
    project.close()


@pytest.mark.parametrize("src_lang", ["ja-JP", "zh-CN"])
def test_export_replay_matches_import_segmentation(tmp_path, src_lang):
    """同一源文件 + 同一源语言：导出重解析必须与导入的分段逐段一致（seq 可重放）。

    这是本修复最容易踩的坑：只改导入、忘了给导出重解析传 src_lang，
    两侧 token 化不一致 → 段文本/段数不同 → 译文按 seq 错位。
    """
    project = Project.create(tmp_path / "q", name="q", src_lang=src_lang,
                             tgt_lang="en-US", provider_id="mock")
    f = tmp_path / "s.md"
    f.write_text(f"# 序章\n\n{ZH_LINE}\n\n{JA_LINE}\n", encoding="utf-8")
    ImportService(project).import_files([f])

    doc = project.db.list_documents()[0]
    segs = project.db.list_segments(doc_id=doc["id"], translatable=True)
    replayed = get_adapter("md").parse(
        project.source_dir() / Path(doc["path"]).name,
        opts={"ruby_loose": project.ruby_loose, "src_lang": project.src_lang})
    again = [b for b in replayed.blocks if b.translatable]
    assert [b.seq for b in again] == [s["seq"] for s in segs]
    assert [b.text for b in again] == [s["src_text"] for s in segs]
    project.close()
