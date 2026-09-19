"""四格式振假名黄金文件测试（增补设计 §1.9）：检测 → 三策略渲染往返。"""
from pathlib import Path

from adapters import get_adapter


def make_md(tmp_path, name="r.md"):
    p = tmp_path / name
    p.write_text("# 序章\n\n僕は北瀬一廣《かずひろ》。`code9` と魔法《まほう》。\n",
                 encoding="utf-8")
    return p


def test_md_ruby_parse_and_three_policies(tmp_path):
    p = make_md(tmp_path)
    ad = get_adapter("md")
    model = ad.parse(p)
    para = next(b for b in model.blocks if b.translatable and "一廣" in b.text)
    assert "{r0}" in para.text and "一廣《" not in para.text
    assert "{r1}" in para.text or "まほう" not in para.text
    entries = {e["token"]: e for e in para.meta["ruby"]}
    assert entries["{r0}"]["rt"] == "かずひろ"
    assert "`code9`" not in para.text  # 行内代码仍是保护占位符

    translations = {b.seq: f"T:{b.text}" for b in model.blocks if b.translatable}
    # keep：token 保留 → 原假名还原
    out = tmp_path / "keep.md"
    ad.render(out, model, translations, "target")
    text = out.read_text(encoding="utf-8")
    assert "T:僕は北瀬一廣《かずひろ》。" in text and "《まほう》" in text
    # translate：注音用译注音
    seq = para.seq
    out2 = tmp_path / "trans.md"
    ad.render(out2, model, translations, "target",
              ruby_maps={seq: {"{r0}": "卡兹希罗", "{r1}": "玛霍", "{r2}": "玛霍"}})
    t2 = out2.read_text(encoding="utf-8")
    assert "《卡兹希罗》" in t2 and "《玛霍》" in t2
    # drop：译文本身无 token → 无注音
    drop_trans = {b.seq: f"T:{b.text}".replace("{r0}", "").replace("{r1}", "")
                          .replace("{r2}", "")
                  for b in model.blocks if b.translatable}
    out3 = tmp_path / "drop.md"
    ad.render(out3, model, drop_trans, "target")
    t3 = out3.read_text(encoding="utf-8")
    assert "《" not in t3 and "一廣" in t3


def test_md_inline_html_ruby_is_tokenized(tmp_path):
    """markdown 里内联的 `<ruby>` 也要被识别为振假名构造。

    回归：md 适配器此前只认青空文库式《》与全角括号，**完全不认内联 `<ruby>`**，
    标签连同内容被当正文送去翻译 —— drop/keep/translate 三种策略全部失效
    （批文本里没有 {rN}，连策略指令都不会注入）。
    """
    p = tmp_path / "r.md"
    p.write_text(
        "# 序章\n\n彼は<ruby>魔法<rt>まほう</rt></ruby>を使った。\n\n"
        "そして<ruby>剣<rt>けん</rt></ruby>と<ruby>盾<rt>たて</rt></ruby>を得た。\n",
        encoding="utf-8")
    ad = get_adapter("md")
    model = ad.parse(p)

    para = next(b for b in model.blocks if b.translatable and "魔法" in b.text)
    # 基词保留在送翻文本中，注音进 token 槽；<rt> 内容绝不进送翻文本
    assert "魔法{r0}" in para.text, para.text
    assert "<rt>" not in para.text and "まほう" not in para.text
    entries = {e["token"]: e for e in para.meta["ruby"]}
    assert entries["{r0}"]["style"] == "html"
    assert entries["{r0}"]["rt"] == "まほう"

    # 同段两个构造都要有独立槽位
    para2 = next(b for b in model.blocks if b.translatable and "剣" in b.text)
    assert "{r0}" in para2.text and "{r1}" in para2.text, para2.text

    translations = {b.seq: b.text for b in model.blocks if b.translatable}

    # keep：保持源文件的 HTML 形态，注音用原假名
    out = tmp_path / "keep.md"
    ad.render(out, model, translations, "target")
    text = out.read_text(encoding="utf-8")
    assert "<ruby>魔法<rt>まほう</rt></ruby>" in text, text
    # 同段连续两个 <ruby> 都能正确锚定（前一个的 </ruby> 不能挡住后一个）
    assert "<ruby>剣<rt>けん</rt></ruby>" in text and "<ruby>盾<rt>たて</rt></ruby>" in text, text

    # translate：注音换成译注音
    out2 = tmp_path / "trans.md"
    ad.render(out2, model, translations, "target",
              ruby_maps={para.seq: {"{r0}": "mó fǎ"}})
    t2 = out2.read_text(encoding="utf-8")
    assert "<ruby>魔法<rt>mó fǎ</rt></ruby>" in t2, t2

    # drop：译文里没有 token → 不应残留 {rN} 或 ruby 标签
    drop = {b.seq: b.text.replace("{r0}", "").replace("{r1}", "").replace("{r2}", "")
            for b in model.blocks if b.translatable}
    out3 = tmp_path / "drop.md"
    ad.render(out3, model, drop, "target")
    t3 = out3.read_text(encoding="utf-8")
    assert "{r" not in t3 and "<ruby>" not in t3, t3


def test_md_inline_html_ruby_variants(tmp_path):
    """`<rb>` / `<rp>` 变体与带属性的 `<ruby>` 也要能识别。"""
    p = tmp_path / "v.md"
    p.write_text(
        "A <ruby><rb>魔法</rb><rt>まほう</rt></ruby> B\n\n"
        "C <ruby class=\"r\">剣<rp>(</rp><rt>けん</rt><rp>)</rp></ruby> D\n",
        encoding="utf-8")
    ad = get_adapter("md")
    model = ad.parse(p)
    entries = [e for b in model.blocks for e in (b.meta.get("ruby") or [])]
    rts = {e["rt"] for e in entries}
    assert rts == {"まほう", "けん"}, rts
    assert all("<rt>" not in b.text for b in model.blocks if b.translatable)


def test_txt_loose_ruby(tmp_path):
    p = tmp_path / "r.txt"
    p.write_text("魔法（まほう）が使える。\n", encoding="utf-8")
    ad = get_adapter("txt")
    m1 = ad.parse(p)
    assert "{r0}" not in next(b.text for b in m1.blocks if b.translatable)
    m2 = ad.parse(p, opts={"ruby_loose": True})
    b = next(b for b in m2.blocks if b.translatable)
    assert b.text == "魔法{r0}が使える。"
    out = tmp_path / "o.txt"
    ad.render(out, m2, {b.seq: f"T:{b.text}"}, "target")
    assert "T:魔法（まほう）が使える。" in out.read_text(encoding="utf-8")


def test_html_ruby_native_restore(tmp_path):
    p = tmp_path / "r.html"
    p.write_text(
        "<html><body><p>彼の名は<ruby>一廣<rt>かずひろ</rt></ruby>だ。</p></body></html>",
        encoding="utf-8")
    ad = get_adapter("html")
    model = ad.parse(p)
    b = next(b for b in model.blocks if b.translatable)
    assert "一廣{r0}" in b.text and b.meta["ruby"][0]["rt"] == "かずひろ"

    translations = {b.seq: f"T:{b.text}" for b in model.blocks if b.translatable}
    out = tmp_path / "keep.html"
    ad.render(out, model, translations, "target")
    text = out.read_text(encoding="utf-8")
    # 锚定译词重建原生结构：按**源基词长度**截短 —— 基词只有 `一廣`，
    # 前面的 `彼の名は` 必须留在 <ruby> 之外（未截短会整串进基词）
    assert "<ruby>一廣<rt>かずひろ</rt></ruby>" in text
    assert "T:彼の名は<ruby>" in text

    out2 = tmp_path / "trans.html"
    ad.render(out2, model, translations, "target", ruby_maps={b.seq: {"{r0}": "卡兹希罗"}})
    t2 = out2.read_text(encoding="utf-8")
    assert "<rt>卡兹希罗</rt>" in t2


def _add_ruby_run(para, base, rt):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    para._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:ruby>'
        f'<w:rubyPr><w:rubyAlign w:val="distributeSpace"/></w:rubyPr>'
        f'<w:rt><w:r><w:t>{rt}</w:t></w:r></w:rt>'
        f'<w:rubyBase><w:r><w:t>{base}</w:t></w:r></w:rubyBase>'
        f'</w:ruby></w:r>'))


def test_docx_ruby_parse_and_native_restore(tmp_path):
    from docx import Document
    src = tmp_path / "r.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("彼の名は")
    _add_ruby_run(para, "一廣", "かずひろ")
    para.add_run("だ。")
    doc.save(str(src))

    ad = get_adapter("docx")
    model = ad.parse(src)
    b = next(b for b in model.blocks if b.translatable)
    assert b.text == "彼の名は一廣{r0}だ。"
    assert b.meta["ruby"][0]["rt"] == "かずひろ"

    translations = {b.seq: f"T:{b.text}" for b in model.blocks if b.translatable}
    out = tmp_path / "keep.docx"
    ad.render(out, model, translations, "target")
    xml = Document(str(out)).paragraphs[0]._p.xml
    assert "<w:rt>" in xml and "かずひろ" in xml          # 原生 w:ruby 保留（注音原样）
    assert "<w:rubyBase>" in xml and "一廣" in xml        # 尾串锚点进入 rubyBase

    out2 = tmp_path / "trans.docx"
    ad.render(out2, model, translations, "target",
              ruby_maps={b.seq: {"{r0}": "卡兹希罗"}})
    xml2 = Document(str(out2)).paragraphs[0]._p.xml
    assert "卡兹希罗" in xml2
