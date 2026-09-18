from pathlib import Path

from adapters import get_adapter


def make_docx(path: Path) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading("第一章 试炼", level=1)
    p = doc.add_paragraph()
    r = p.add_run("林凡握着")
    r.bold = True
    p.add_run("一块灵石。")
    doc.add_paragraph("第二段正文。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "灵石"
    table.cell(1, 1).text = "三块"
    doc.save(str(path))


def test_docx_parse(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src)
    model = get_adapter("docx").parse(src)
    translatable = [b for b in model.blocks if b.translatable]
    texts = [b.text for b in translatable]
    assert "第一章 试炼" in texts
    assert "林凡握着一块灵石。" in texts
    assert "单元格" not in " ".join(texts)  # 表格单元格是独立段
    assert any(t == "名称" for t in texts) and any(t == "灵石" for t in texts)
    heading = next(b for b in translatable if b.text == "第一章 试炼")
    assert heading.is_heading


def test_docx_render_target_keeps_style(tmp_path):
    from docx import Document
    src = tmp_path / "a.docx"
    make_docx(src)
    ad = get_adapter("docx")
    model = ad.parse(src)
    translations = {b.seq: f"T:{b.text}" for b in model.blocks if b.translatable}
    out = tmp_path / "out.docx"
    ad.render(out, model, translations, "target")
    doc = Document(str(out))
    paras = [p.text for p in doc.paragraphs]
    assert paras[0] == "T:第一章 试炼"
    assert paras[1] == "T:林凡握着一块灵石。"   # 整段套用主样式（边界已知）
    # 标题样式保留
    assert doc.paragraphs[0].style.name.startswith("Heading") or \
        doc.paragraphs[0].style.name.startswith("标题")
    table = doc.tables[0]
    assert table.cell(1, 0).text == "T:灵石"


def test_docx_render_bilingual(tmp_path):
    from docx import Document
    src = tmp_path / "a.docx"
    make_docx(src)
    ad = get_adapter("docx")
    model = ad.parse(src)
    translations = {b.seq: f"T:{b.text}" for b in model.blocks if b.translatable}
    out = tmp_path / "bi.docx"
    ad.render(out, model, translations, "bi_inter")
    doc = Document(str(out))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "第一章 试炼" in paras and "T:第一章 试炼" in paras

    out2 = tmp_path / "bit.docx"
    ad.render(out2, model, translations, "bi_table")
    doc2 = Document(str(out2))
    assert doc2.tables[0].cell(0, 0).text == "原文"
    assert doc2.tables[0].cell(1, 0).text == "第一章 试炼"
