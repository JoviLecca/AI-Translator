"""M2 验收测试（设计 §11 M2）：docx 小说项目端到端。

导入 docx（标题/正文/表格）→ 打标签与风格 → 术语自动归纳（两阶段）→ 审核 →
批量翻译（术语一致）→ TM 跨文件复用 → 校对编辑/确认 → 导出 docx（样式保留）+ html。
"""
import json
from pathlib import Path

from docx import Document

from core.pipeline import ExportService, ImportService, ReviewService, TranslationService
from core.project import Project
from core.term_induction import InductionService
from llm.mock import MockProvider


def make_chapter(path: Path, title: str) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("林凡走进宗门大殿，掌门递给他一块灵石。")
    doc.add_paragraph("灵石可以兑换丹药，宗门每月发放俸禄。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "灵石"
    table.cell(1, 1).text = "三块"
    doc.save(str(path))


def induction_translator(user: str) -> str:
    return json.dumps({"terms": [
        {"src": "灵石", "candidates": ["spirit stone"], "note": "修炼货币"},
        {"src": "宗门", "candidates": ["sect"], "note": "组织"},
        {"src": "林凡", "candidates": ["Lin Fan"], "note": "主角"},
    ]}, ensure_ascii=False)


def translate(text: str) -> str:
    for src, tgt in (("灵石", "spirit stone"), ("宗门", "sect"), ("林凡", "Lin Fan")):
        text = text.replace(src, tgt)
    return "T:" + text


def test_m2_docx_end_to_end(tmp_path):
    # 1) 项目 + docx 章节
    project = Project.create(tmp_path / "novel", name="M2验收", src_lang="zh-CN",
                             tgt_lang="en-US", style_preset="literary", provider_id="mock")
    ch1, ch2 = tmp_path / "ch1.docx", tmp_path / "ch2.docx"
    make_chapter(ch1, "第一章 试炼")
    make_chapter(ch2, "第二章 交锋")
    imported, errs = ImportService(project).import_files([ch1, ch2])
    assert not errs and len(imported) == 2

    # 2) 文件标签（FR-04）
    ImportService(project).set_tags(imported[0]["doc_id"],
                                    {"genre": "奇幻小说", "style": "网文口语"})

    # 3) 术语自动归纳 → 审核 → 入库
    inducer = InductionService(project, MockProvider(translator=induction_translator))
    res = inducer.run()
    assert res["docs"] == 2
    assert {c["src"] for c in res["candidates"]} == {"灵石", "宗门", "林凡"}
    inducer.approve([("灵石", ["spirit stone"], "货币"), ("宗门", ["sect"], "组织"),
                     ("林凡", ["Lin Fan"], "主角")])
    assert project.glossary.get("灵石").candidates == ["spirit stone"]
    # 增量：归纳后新文件只处理新文件
    ch3 = tmp_path / "ch3.docx"
    make_chapter(ch3, "第三章 归途")
    ImportService(project).import_files([ch3])
    res2 = inducer.run()
    assert res2["docs"] == 1

    # 4) 预检 + 批量翻译（术语强制一致）
    svc = TranslationService(project, {"concurrency": 4})
    pre = svc.precheck()
    assert pre["pending"] > 0 and pre["glossary_terms"] == 3
    handle = svc.start(provider=MockProvider(translator=translate))
    handle.join(timeout=60)
    assert handle.result["failed"] == 0, handle.result

    segs = project.db.list_segments(translatable=True)
    assert all(s["tgt_text"] for s in segs)
    for s in segs:
        if "灵石" in s["src_text"]:
            assert "spirit stone" in s["tgt_text"]

    # 5) TM 跨文件复用：ch2/ch3 与 ch1 有相同段落 → 运行内去重回填
    assert handle.result["tm"] + handle.result.get("dedup", 0) > 0

    # 6) 校对：修改 + 确认
    review = ReviewService(project)
    row = next(r for r in review.segments() if "俸禄" in r["src"])
    review.edit(row["id"], "T:sect pays a monthly stipend.")
    review.confirm(row["id"])
    assert next(r for r in review.segments() if r["id"] == row["id"])["status"] == "confirmed"

    # 7) 导出 docx（样式保留）+ html 章节端到端
    doc_ids = [d["id"] for d in project.db.list_documents()]
    results = ExportService(project).export(doc_ids, mode="target")
    assert all(r["ok"] for r in results), results
    out1 = project.target_dir() / "ch1.en.docx"
    assert out1.exists()
    reread = Document(str(out1))
    assert reread.paragraphs[0].text == "T:第一章 试炼"
    assert reread.paragraphs[0].style.name.startswith(("Heading", "标题"))
    assert "T:spirit stone" in reread.paragraphs[1].text or \
        "spirit stone" in reread.paragraphs[1].text
    assert reread.tables[0].cell(1, 0).text == "T:spirit stone"

    # html 文件走同一管线
    html_file = tmp_path / "ch1.html"
    html_file.write_text(
        "<html><body><h1>试炼</h1><p>林凡握着灵石。</p></body></html>", encoding="utf-8")
    ImportService(project).import_files([html_file])
    h2 = svc.start(provider=MockProvider(translator=translate))
    h2.join(timeout=60)
    assert h2.result["failed"] == 0
    html_doc = next(d for d in project.db.list_documents() if d["format"] == "html")
    r = ExportService(project).export([html_doc["id"]], mode="target")[0]
    assert r["ok"]
    assert "spirit stone" in Path(r["out"]).read_text(encoding="utf-8")

    # 8) 源文件变更防护：改 source 后导出应告警
    (project.source_dir() / "ch1.html").write_text(
        "<html><body><h1>试炼</h1><p>内容已改。</p></body></html>", encoding="utf-8")
    r2 = ExportService(project).export([html_doc["id"]], mode="target", force=False)[0]
    assert not r2["ok"] and any("源文件" in w for w in r2["warnings"])
    r3 = ExportService(project).export([html_doc["id"]], mode="target", force=True)[0]
    assert r3["ok"]

    project.close()
